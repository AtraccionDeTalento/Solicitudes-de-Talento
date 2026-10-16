# Extrae las DP del formato Excel de USIL (hoja MODELO) a tools/dps.local.json y las listas de «Desplegables» a dp-listas.js.
# Uso: python -I tools/extraer-dps.py "C:/ruta/DP1.xlsx" "C:/ruta/DP2.xlsx" ...   (dps.local.json queda fuera de git por .gitignore)
import json, sys, re, os, warnings
warnings.filterwarnings("ignore")
import openpyxl
AQUI = os.path.dirname(os.path.abspath(__file__))
t = lambda v: "" if v is None else re.sub(r"\s+", " ", str(v)).strip()

def extraer(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True); w = wb["MODELO"]; c = lambda a: t(w[a].value)
    marcas = lambda filas, col_n, col_v: [{"con": c(f"{col_n}{r}"), "tipo": c(f"{col_v}{r}")} for r in filas if c(f"{col_n}{r}")]
    niveles = lambda filas, cn, cv: [{"nombre": c(f"{cn}{r}"), "nivel": c(f"{cv}{r}")} for r in filas if c(f"{cn}{r}")]
    auto = next((c(f"B{r}")[0] for r in range(43, 47) if c(f"H{r}")), "")
    firmas = []
    for r in range(80, w.max_row):
        if c(f"B{r}").startswith("Firma de:"):
            firmas = [{"cargo": re.sub(r"^(Firma de|Revisado por):\s*", "", c(f"{col}{r}")), "nombre": c(f"{col}{r+1}")} for col in "BEH"]
            fila_firmas = r; break
    else: fila_firmas = w.max_row
    obs = " ".join(c(f"B{r}") for r in range(82, fila_firmas) if c(f"B{r}") and not c(f"B{r}").startswith("#"))
    funciones = [c(f"C{r}") for r in range(21, 31) if c(f"C{r}")]
    return {
      "titulo": c("C6"), "departamento": c("C8"),
      "datos": {
        "titulo": c("C6"), "seccion": c("C7"), "area": c("G7"), "departamento": c("C8"), "grupoOcupacional": c("G8"), "categoria": c("C9"),
        "unidadNegocio": c("G9"), "centroCosto": c("C10"), "gradoSalarial": c("G10"), "reportaA": c("C11"), "supervisaA": c("G11"),
        "mision": c("B16"), "funciones": funciones,
        "relInternas": marcas(range(34, 41), "B", "D"), "relExternas": marcas(range(34, 41), "F", "H"),
        "autonomia": auto,
        "educacion": {"grado": c("F49"), "profesion": c("F50"), "colegiatura": c("F51"), "especializaciones": c("B53"), "experienciaAnios": c("B55"), "experienciaMinPuesto": c("F55")},
        "idiomas": niveles(range(58, 61), "B", "H"), "ofimatica": niveles(range(62, 67), "B", "H"),
        "competencias": {"transversales": niveles(range(69, 74), "B", "D"), "lideres": niveles(range(69, 71), "F", "H"), "adicionales": niveles(range(72, 76), "F", "H")},
        "equipos": {"laptop": c("H76"), "pc": c("H77"), "movil": c("H78"), "otros": c("H79"), "cajaChica": c("H80")},
        "observaciones": obs, "firmas": firmas }}

def listas(ruta):
    w = openpyxl.load_workbook(ruta, data_only=True)["Desplegables"]; out = {}
    for col in w.iter_cols(max_row=w.max_row):
        v = [t(x.value) for x in col if t(x.value)]
        if v: out[v[0]] = v[1:]
    return out

if __name__ == "__main__":
    archivos = sys.argv[1:]
    if not archivos: sys.exit(__doc__)
    dps = []
    for i, f in enumerate(archivos, 1):
        d = extraer(f); d["codigo"] = "DP-%03d" % i; dps.append(d)
        print(d["codigo"], d["titulo"], "|", d["departamento"], "|", len(d["datos"]["funciones"]), "funciones")
    json.dump(dps, open(os.path.join(AQUI, "dps.local.json"), "w", encoding="utf8"), ensure_ascii=False, indent=1)
    L = listas(archivos[0])
    pick = lambda k: [x for x in L.get(k, []) if x != "Total general"]
    LISTAS = {"seccion": pick("Sección"), "area": pick("Área"), "grupoOcupacional": pick("Grupo Ocupacional"), "categoria": pick("Categoría de Puesto"),
      "unidadNegocio": pick("Unidad de Negocio"), "idiomas": pick("Idiomas"), "nivelIdioma": pick("Idioma"), "programas": pick("Programas"), "nivelPrograma": pick("Nivel Programas"),
      "nivelCompetencia": pick("Nivel Competencias"), "colegiatura": pick("Colegiatura"), "experiencia": pick("Experiencia Mínima Requerida"),
      "grado": ["Bachiller", "Maestro", "Doctor", "No aplica"], "experienciaMinPuesto": ["Hasta 1 año", "1 a 3 años", "3 a 5 años", "5 años a más", "4 años a más (Docente)", "No aplica"],
      "relacion": ["Enviar / Recibir Información", "Coordinar", "Negociar", "Atención", "Productos y servicios", "Fidelización", "Eventos", "No aplica"]}
    open(os.path.join(AQUI, "..", "dp-listas.js"), "w", encoding="utf8").write("// Listas desplegables de la DP (hoja Desplegables del formato de USIL). Generado por tools/extraer-dps.py.\nexport const LISTAS = " + json.dumps(LISTAS, ensure_ascii=False, indent=1) + ";\n")
    print({k: len(v) for k, v in LISTAS.items()})
